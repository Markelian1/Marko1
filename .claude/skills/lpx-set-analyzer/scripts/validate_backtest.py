#!/usr/bin/env python3
"""Validate an MT5 backtest XLSX (Spanish or English report) for a chosen set.
Usage: python3 validate_backtest.py file.xlsx [--expected-pf 1.46] [--deploy-risk 0.3] [--target 5] [--set winner.txt]
       [--pf-tol 0.05] [--window-months 2] [--out-dir .]"""
import os, re, sys, argparse, datetime as dt, calendar
import numpy as np, pandas as pd
from openpyxl import load_workbook

def num(v):
    if isinstance(v,bool): return None
    if isinstance(v,(int,float)): return float(v)
    if isinstance(v,str):
        m=re.match(r'^\s*(-?[\d\s ]*\.?\d+)',v)
        if m:
            try: return float(re.sub(r'[\s ]','',m.group(1)))
            except ValueError: return None
    return None

def label_value(rows,labels):
    """First number to the right of a cell starting with one of the labels."""
    for row in rows:
        for j,c in enumerate(row):
            if isinstance(c,str) and any(c.strip().startswith(l) for l in labels):
                for v in row[j+1:]:
                    x=num(v)
                    if x is not None: return x
    return None

def same(a,b):
    fa,fb=num(a),num(b)
    if fa is not None and fb is not None: return abs(fa-fb)<1e-9
    return str(a).strip().lower()==str(b).strip().lower()

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('--expected-pf',type=float,default=None)
    ap.add_argument('--deploy-risk',type=float,default=None); ap.add_argument('--target',type=float,default=5.0)
    ap.add_argument('--set',default=None,help='key=value file (set_analyzer_winner.txt) to check the report inputs against')
    ap.add_argument('--risk-key',default='RiskPercent'); ap.add_argument('--pf-tol',type=float,default=0.05)
    ap.add_argument('--window-months',type=int,default=2); ap.add_argument('--out-dir',default='.')
    a=ap.parse_args()
    wb=load_workbook(a.xlsx,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    rows=list(ws.iter_rows(values_only=True))
    # deals table header
    dh=None
    for i,row in enumerate(rows):
        vals=[str(c).strip() for c in row if c is not None]
        if 'Balance' in vals and ('Beneficio' in vals or 'Profit' in vals) and ('Dirección' in vals or 'Direction' in vals): dh=i; break
    if dh is None: sys.exit("!! Deals table (Beneficio/Balance/Dirección) not found — is this the MT5 backtest report?")
    head=rows[:dh]
    # inputs & summary
    inputs={}
    for row in head:
        for c in row:
            m=re.match(r'^\s*([A-Za-z_]\w*)=(.*)$',c) if isinstance(c,str) else None
            if m: inputs[m.group(1)]=m.group(2).strip()
    for k,v in inputs.items(): print(f'CFG: {k}={v}')
    keys=['Beneficio Neto','Factor de Beneficio','Reducción máxima del balance','Ratio de Sharpe','Factor de Recuperación',
          'Total de operaciones','rentables','cortas','largas','Total Net Profit','Profit Factor','Sharpe Ratio','Recovery Factor','Total Trades']
    for row in head:
        j=' | '.join(str(c) for c in row if c is not None)
        if any(k in j for k in keys): print('SUMMARY:',j[:140])
    native_risk=num(inputs.get(a.risk_key)); deposit=label_value(head,['Depósito inicial','Initial Deposit'])
    report_pf=label_value(head,['Factor de Beneficio','Profit Factor'])
    print(f"Native {a.risk_key}: {native_risk}" if native_risk else f"!! {a.risk_key} not found in report inputs — no risk scaling possible")
    if a.set:
        exp={}
        for ln in open(a.set,encoding='utf-8'):
            if '=' in ln and not ln.lstrip().startswith(';'):
                k,v=ln.split('=',1); exp[k.strip()]=v.split('||')[0].strip()
        bad=[(k,v,inputs.get(k)) for k,v in exp.items() if k!=a.risk_key and (k not in inputs or not same(v,inputs[k]))]
        if bad:
            print("!! PARAM MISMATCH vs recommended set — wrong set loaded, do not deploy:")
            for k,v,got in bad: print(f"   {k}: expected {v} | report {got if got is not None else '(missing)'}")
        else: print(f"PARAMS MATCH recommended set ({len(exp)} keys; {a.risk_key} may differ — returns normalize)")
    # deals: column positions from the header row, falling back to the standard MT5 layout
    hdr=[str(c).strip() if c is not None else '' for c in rows[dh]]
    def col(names,default):
        for n in names:
            if n in hdr: return hdr.index(n)
        return default
    iT=col(['Tipo','Type'],3); iD=col(['Dirección','Direction'],4); iP=col(['Beneficio','Profit'],10); iB=col(['Balance'],11)
    closes=[]; daylink=[]; cur=None; shorts=0; longs=0; prev=None; pend=0.0; first_bal=None
    for r in rows[dh+1:]:
        t=r[0]
        if isinstance(t,str):
            try: t=dt.datetime.strptime(t.strip(),'%Y.%m.%d %H:%M:%S')
            except ValueError: t=None
        if not isinstance(t,dt.datetime): continue
        typ=str(r[iT]).strip().lower() if r[iT] is not None else ''; dr=str(r[iD]).strip().lower() if r[iD] is not None else ''
        if dr=='in':
            cur=t
            if typ=='buy': longs+=1
            elif typ=='sell': shorts+=1
        elif dr in ('out','out by','in/out') and isinstance(r[iP],(int,float)) and cur is not None:
            daylink.append((cur,t,r[iP]))
        bal=num(r[iB])
        if bal is None: continue
        if prev is None: prev=deposit if deposit else bal; first_bal=prev
        delta=bal-prev; prev=bal
        if typ in ('balance','credit'): continue      # deposits/withdrawals are not trading P&L
        if dr=='in': pend+=delta; continue             # entry commission rides with the next close
        pnl=delta+pend; pend=0.0
        if abs(pnl)>1e-9: closes.append((t,pnl,bal))
    if not closes: sys.exit("!! No realized closes found in the deals table.")
    real=pd.DataFrame(closes,columns=['time','pnl','balance'])
    real['r']=real.pnl/(real.balance-real.pnl)
    net=(real.balance.iloc[-1]/first_bal-1)*100
    gross_w=real.pnl[real.pnl>0].sum(); gross_l=-real.pnl[real.pnl<0].sum()
    pf=gross_w/gross_l if gross_l>0 else float('inf')
    print(f"\nInitial deposit: {first_bal:,.2f} | Closes: {len(real)} | Net: {net:.1f}% | report PF: {report_pf} | close-level PF: {pf:.3f} | close win% {(real.pnl>0).mean()*100:.1f}")
    if a.expected_pf:
        use=report_pf if report_pf else pf; diff=use-a.expected_pf
        print(f"PF MATCH vs optimization {a.expected_pf}: {use:.3f} ({diff:+.3f}, tol ±{a.pf_tol}) -> "
              f"{'OK' if abs(diff)<=a.pf_tol else '!! MISMATCH — wrong set or data issue, do not deploy'}")
    print(f"DIRECTION: longs {longs} shorts {shorts}" + ("  !! LONG-ONLY (regime-dependent)" if shorts==0 else "") + ("  !! SHORT-ONLY" if longs==0 else ""))
    # day-of-week attribution
    if daylink:
        L=pd.DataFrame(daylink,columns=['tin','tout','pnl']); L['dow']=L.tin.dt.weekday
        names={0:'Lun',1:'Mar',2:'Mié',3:'Jue',4:'Vie',5:'Sáb',6:'Dom'}
        g=L.groupby('dow').agg(n=('pnl','size'),sum_usd=('pnl','sum'),win=('pnl',lambda x:round((x>0).mean()*100,1)))
        tot=L.pnl.sum(); g['pct_profit']=(g.sum_usd/tot*100).round(1) if tot else np.nan
        g.index=[names.get(i,i) for i in g.index]
        print("\nP&L by ENTRY day:\n"+g.round(1).to_string())
        if (g['pct_profit']>70).any(): print("!! DAY CONCENTRATION: one entry day carries >70% of profit — fragile edge, check session")
        overnight=(L.tout.dt.date>L.tin.dt.date).sum(); wknd=((L.tout.dt.weekday<L.tin.dt.weekday)|((L.tout-L.tin).dt.days>=2)).sum()
        print(f"Overnight holds: {overnight}/{len(L)} | weekend-crossing: ~{wknd}  (Master: weekend prohibited -> re-baseline if material)")
    # rolling windows & monthly table at deploy risk
    mult=(a.deploy_risk/native_risk) if (a.deploy_risk and native_risk) else 1.0
    print(f"\nScaling to deploy risk {a.deploy_risk}% (native {native_risk}%): x{mult:.2f}" if mult!=1.0 else "\n(no deploy-risk scaling)")
    real['rs']=real.r*mult
    W=a.window_months
    def plusW(tt):
        m2=tt.month-1+W; y=tt.year+m2//12; m2=m2%12+1
        return pd.Timestamp(y,m2,min(tt.day,calendar.monthrange(y,m2)[1]))
    t0,tN=real.time.iloc[0],real.time.iloc[-1]; w=[]
    for s in pd.date_range(t0.normalize(),tN.normalize(),freq='D'):
        e=plusW(s)
        if e>tN: break
        rs=real.rs.values[(real.time>=s).values&(real.time<e).values]
        w.append((np.prod(1+rs)-1)*100 if len(rs) else 0.0)
    eq=np.cumprod(1+real.rs.values); pk=np.maximum.accumulate(np.r_[1.0,eq])[1:]
    if w:
        w=np.array(w)
        print(f"{W}-mo windows: n={len(w)} median {np.median(w):.2f}% P(>0) {(w>0).mean():.3f} P(>={a.target:g}%) {(w>=a.target).mean():.3f} "
              f"worst {w.min():.2f}% worst5% {np.percentile(w,5):.2f}%")
    else: print(f"!! Backtest shorter than {W} months — no rolling windows")
    worst=real.rs.min()*100; best=real.rs.max()*100
    print(f"MaxDD @deploy: {((eq-pk)/pk).min()*100:.2f}% | worst single close: {worst:.3f}% | biggest single win: {best:.3f}% "
          f"({best/a.target*100:.0f}% of {a.target:g}% target)")
    if worst<=-1.2: print("!! Worst close breaches 1.2% (Striking limit) at deploy risk")
    if best>=0.6*a.target: print("!! Single win >= 60% of the phase target — consistency-rule risk")
    mo=real.groupby(real.time.dt.to_period('M')).agg(closes=('rs','size'),ret_pct=('rs',lambda x:(np.prod(1+x)-1)*100))
    print("\nMonthly @deploy:\n"+mo.round(2).to_string())
    os.makedirs(a.out_dir,exist_ok=True); path=os.path.join(a.out_dir,'set_analyzer_backtest.csv')
    real[['time','r','balance']].to_csv(path,index=False)
    print(f"\nStream (r at native risk) saved to {path} (for portfolio combination via lpx-portfolio-manager)")
if __name__=='__main__': main()
