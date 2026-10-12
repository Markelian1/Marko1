"""Gjeneron skedarin Raporti_i_Analizave.xlsx (ditar tregtimi / trading journal).

Përdorimi:
    python3 build_raporti.py            -> Raporti_i_Analizave.xlsx (me 1 rresht shembull)
    python3 build_raporti.py --demo OUT -> skedar demonstrues me trade-t e raportit origjinal
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import AreaChart, DoughnutChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.chart.axis import ChartLines
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun)
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- konstante
FONT = "Arial"
FIRST = 12            # rreshti i parë i trade-ve
LAST = 511            # rreshti i fundit (500 trade)
MAIN = "Raporti"
SET = "Cilësimet"
SETQ = f"'{SET}'"

BG = "0B0B0B"          # sfondi i përgjithshëm
ROW_BG = "161616"      # rreshtat e tabelës
INPUT_BG = "1D1D1D"    # qelizat që plotësohen
HEAD_BG = "1A2238"     # header-at
LABEL_BG = "17233B"
VALUE_BG = "222222"
GRID = "2E2E2E"
WHITE = "FFFFFF"
GREY = "9AA0A6"
BLUE = "8AB4F8"
GREEN = "22C55E"
GREEN_DARK = "14532D"
GREEN_LINE = "39E75F"
RED = "EF4444"

STATUSES = [  # statusi, kategoria, × R:R, + R fikse
    ("TP", "WIN", 1, 0),
    ("CLOSE HALF - TP", "WIN", 0.75, 0),
    ("CLOSE HALF - BE", "BE", 0.25, 0),
    ("BE", "BE", 0, 0),
    ("LOSS", "LOSS", 0, -1),
]
STATUS_STYLE = {  # mbushja, ngjyra e tekstit
    "TP": ("15803D", WHITE),
    "CLOSE HALF - TP": ("57BB8A", "0B0B0B"),
    "CLOSE HALF - BE": ("B7E1CD", "0B0B0B"),
    "BE": ("E8EAED", "3C4043"),
    "LOSS": ("C62828", WHITE),
}
SETUPS = ["IMB", "CHILL", "BOS", "CHOCH", "FVG", "OB", "BREAKER", "SWEEP"]
PAIRS = ["EUR/USD", "GBP/USD", "USD/JPY", "USD/CHF", "USD/CAD", "AUD/USD", "NZD/USD",
         "EUR/JPY", "EUR/GBP", "EUR/CAD", "EUR/AUD", "EUR/CHF", "GBP/JPY", "GBP/AUD",
         "AUD/JPY", "AUD/CHF", "CAD/JPY", "XAU/USD", "US30", "NAS100"]

# Trade-t e dukshme në raportin origjinal (përdoren vetëm për --demo)
DEMO = [
    ("2026-01-07", "EUR/JPY", "IMB", "CLOSE HALF - BE"), ("2026-01-07", "EUR/CAD", "IMB", "CLOSE HALF - BE"),
    ("2026-01-13", "EUR/CAD", "CHILL", "BE"), ("2026-01-13", "EUR/USD", "CHILL", "CLOSE HALF - BE"),
    ("2026-01-14", "GBP/AUD", "IMB", "BE"), ("2026-01-16", "US30", "CHILL", "CLOSE HALF - TP"),
    ("2026-01-20", "AUD/CHF", "IMB", "LOSS"), ("2026-01-20", "EUR/AUD", "IMB", "CLOSE HALF - TP"),
    ("2026-01-21", "EUR/CAD", "IMB", "CLOSE HALF - BE"), ("2026-01-21", "EUR/USD", "IMB", "BE"),
    ("2026-01-21", "USD/CHF", "IMB", "CLOSE HALF - BE"), ("2026-02-10", "AUD/USD", "CHILL", "BE"),
    ("2026-02-16", "EUR/JPY", "CHILL", "CLOSE HALF - TP"), ("2026-02-20", "EUR/GBP", "CHILL", "CLOSE HALF - BE"),
    ("2026-02-25", "EUR/JPY", "IMB", "BE"), ("2026-03-04", "GBP/JPY", "IMB", "CLOSE HALF - BE"),
    ("2026-03-18", "USD/JPY", "IMB", "BE"), ("2026-03-25", "EUR/CAD", "CHILL", "BE"),
    ("2026-04-02", "CAD/JPY", "CHILL", "BE"), ("2026-04-14", "EUR/GBP", "IMB", "CLOSE HALF - BE"),
    ("2026-04-30", "GBP/USD", "CHILL", "CLOSE HALF - TP"), ("2026-05-05", "AUD/USD", "CHILL", "CLOSE HALF - TP"),
]


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def font(color=WHITE, size=10, bold=False, italic=False):
    return Font(name=FONT, color=color, size=size, bold=bold, italic=italic)


CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
thin = Side(style="thin", color=GRID)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
white_side = Side(style="medium", color="DADCE0")
WBOX = Border(left=white_side, right=white_side, top=white_side, bottom=white_side)


# ---------------------------------------------------------------- grafikët
def char_props(color, size, bold=False):
    return CharacterProperties(sz=size, b=bold, solidFill=color, latin=DFont(typeface=FONT))


def chart_title(text, color=WHITE, size=1400):
    cp = char_props(color, size, bold=True)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def axis_text(color=GREY, size=900):
    cp = char_props(color, size, bold=True)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def dark_background(chart):
    chart.graphical_properties = GraphicalProperties(solidFill=BG, ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(solidFill=BG, ln=LineProperties(noFill=True))


def balance_chart(ws, last_row, width, height, title_size):
    ch = AreaChart()
    ch.title = chart_title("Balance", size=title_size)
    ch.legend = None
    ch.add_data(Reference(ws, min_col=8, min_row=FIRST, max_row=last_row), titles_from_data=False)
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=GREEN_DARK,
                                                ln=LineProperties(solidFill=GREEN_LINE, w=28575))
    ch.y_axis.number_format = "0%"
    ch.y_axis.delete = False
    ch.y_axis.txPr = axis_text()
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="3C4043", w=9525)))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.delete = True
    ch.display_blanks = "gap"
    dark_background(ch)
    ch.width, ch.height = width, height
    return ch


def winrate_chart(ws):
    ch = DoughnutChart(holeSize=55)
    ch.legend = None
    ch.add_data(Reference(ws, min_col=3, min_row=5, max_row=6), titles_from_data=False)
    s = ch.series[0]
    for idx, color in ((0, GREEN), (1, RED)):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=BG, w=12700))
        s.dPt.append(pt)
    dark_background(ch)
    ch.layout = Layout(manualLayout=ManualLayout(x=0.05, y=0.05, w=0.9, h=0.9))
    ch.width, ch.height = 3.6, 3.7
    return ch


# ---------------------------------------------------------------- fletët
def build_settings(wb, risk):
    ws = wb.create_sheet(SET)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", (20, 12, 10, 12, 4, 14, 4, 12, 4, 60)):
        ws.column_dimensions[col].width = w
    for r in range(1, 40):
        for c in range(1, 11):
            ws.cell(r, c).fill = fill(BG)
            ws.cell(r, c).font = font()

    ws["A1"] = "CILËSIMET"
    ws["A1"].font = font(size=18, bold=True)
    ws.row_dimensions[1].height = 30

    ws["A3"] = "Rreziku për trade"
    ws["A3"].font = font(BLUE, bold=True)
    ws["C3"] = risk
    ws["C3"].number_format = "0.00%"
    ws["C3"].font = font(bold=True)
    ws["C3"].fill = fill(INPUT_BG)
    ws["C3"].border = BOX
    ws["C3"].alignment = CENTER
    ws["C3"].comment = Comment(
        "Vlerë e zgjedhur nga ti (1% si parazgjedhje). Raporti origjinal përdor 23%: "
        "me 23% dhe R:R 1.40 dalin saktë 8.05% / 24.15% / -23%.", "Raporti")

    heads = ("Statusi", "Kategoria", "× R:R", "+ R fikse")
    for i, h in enumerate(heads):
        c = ws.cell(5, 1 + i, h)
        c.font = font(BLUE, bold=True)
        c.fill = fill(HEAD_BG)
        c.alignment = CENTER
        c.border = BOX
    for r, (st, cat, mult, fixed) in enumerate(STATUSES, start=6):
        for i, v in enumerate((st, cat, mult, fixed)):
            c = ws.cell(r, 1 + i, v)
            c.fill = fill(INPUT_BG)
            c.border = BOX
            c.alignment = CENTER
            c.font = font(bold=(i == 0))
            if i >= 2:
                c.number_format = "0.00"
        bg, fg = STATUS_STYLE[st]
        ws.cell(r, 1).fill = fill(bg)
        ws.cell(r, 1).font = font(fg, bold=True)

    cat_dv = DataValidation(type="list", formula1='"WIN,LOSS,BE"', allow_blank=True)
    ws.add_data_validation(cat_dv)
    cat_dv.add("B6:B10")

    notes = [
        "Net = Rreziku × (R:R × shumëzuesi + R fikse)",
        "TP: e gjithë pozita mbyllet te TP  →  1 × R:R",
        "CLOSE HALF - TP: gjysma mbyllet te gjysma e rrugës, gjysma te TP  →  0.75 × R:R",
        "CLOSE HALF - BE: gjysma mbyllet te gjysma e rrugës, pjesa tjetër në BE  →  0.25 × R:R",
        "BE: mbyllet pa fitim e pa humbje  →  0",
        "LOSS: goditet SL  →  -1R",
        "Kategoria vendos ku numërohet trade-i: WIN, LOSS ose BE.",
        "Winrate = WINS / (WINS + LOSS). Trade-t BE nuk ndikojnë në winrate.",
    ]
    for i, t in enumerate(notes):
        ws.cell(5 + i, 10, t).font = font(GREY, size=9, italic=(i > 0))

    ws["F4"] = "Setup"
    ws["H4"] = "Pair"
    for ref in ("F4", "H4"):
        ws[ref].font = font(BLUE, bold=True)
        ws[ref].fill = fill(HEAD_BG)
        ws[ref].alignment = CENTER
    for i, v in enumerate(SETUPS):
        c = ws.cell(5 + i, 6, v)
        c.fill, c.border, c.alignment = fill(INPUT_BG), BOX, CENTER
    for i, v in enumerate(PAIRS):
        c = ws.cell(5 + i, 8, v)
        c.fill, c.border, c.alignment = fill(INPUT_BG), BOX, CENTER

    ws["A13"] = "Qelizat me sfond pak më të çelët (C3, tabela e statuseve, Setup, Pair) mund t'i ndryshosh."
    ws["A13"].font = font(GREY, size=9, italic=True)
    return ws


def build_main(wb, trades, demo):
    ws = wb.active
    ws.title = MAIN
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN

    widths = {"A": 2, "B": 13, "C": 12, "D": 11, "E": 8, "F": 20, "G": 10, "H": 11, "I": 28, "J": 3}
    for col in range(1, 20):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = widths.get(letter, 13)

    for r in range(1, LAST + 4):
        for c in range(1, 20):
            ws.cell(r, c).fill = fill(BG)
            ws.cell(r, c).font = font()

    # --- titulli dhe periudha
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 44
    ws.merge_cells("B2:I2")
    ws["B2"] = "RAPORTI I ANALIZAVE"
    ws["B2"].font = font(size=30, bold=True)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="bottom")
    ws["B3"] = date(2026, 1, 1)
    ws["C3"] = "deri"
    ws["D3"] = date(2026, 12, 31)
    for ref in ("B3", "D3"):
        ws[ref].number_format = "dd mmm yyyy"
        ws[ref].font = font(bold=True)
        ws[ref].fill = fill(INPUT_BG)
        ws[ref].alignment = CENTER
    ws["C3"].font = font(GREY, bold=True)
    ws["C3"].alignment = CENTER
    ws["B3"].comment = Comment("Data e fillimit të periudhës së raportit (mund ta ndryshosh).", "Raporti")
    ws["D3"].comment = Comment("Data e mbarimit të periudhës së raportit (mund ta ndryshosh).", "Raporti")
    ws.row_dimensions[3].height = 20

    # --- statistikat
    rng = lambda col: f"${col}${FIRST}:${col}${LAST}"
    stats = [
        (4, "Total Trades", BLUE, f"=SUM({rng('Z')})"),
        (5, "Total WINS", GREEN, f'=COUNTIFS({rng("Z")},1,{rng("AB")},"WIN")'),
        (6, "Total LOSS", RED, f'=COUNTIFS({rng("Z")},1,{rng("AB")},"LOSS")'),
        (7, "Total BE", WHITE, f'=COUNTIFS({rng("Z")},1,{rng("AB")},"BE")'),
    ]
    for r, label, color, formula in stats:
        ws.row_dimensions[r].height = 21
        lc = ws.cell(r, 2, label)
        lc.font, lc.fill, lc.alignment, lc.border = font(color, 11, True), fill(LABEL_BG), CENTER, BOX
        vc = ws.cell(r, 3, formula)
        vc.font, vc.fill, vc.alignment, vc.border = font(size=11, bold=True), fill(VALUE_BG), CENTER, BOX
        vc.number_format = "0"
    for r in (8, 9):
        ws.row_dimensions[r].height = 21
    ws.merge_cells("B8:B9")
    ws.merge_cells("C8:C9")
    ws["B8"] = "PnL"
    ws["B8"].font = font(size=15, bold=True)
    ws["C8"] = f"=EXP(SUM($AA${FIRST}:$AA${LAST}))-1"
    ws["C8"].font = font(GREEN, 15, True)
    ws["C8"].number_format = "0.00%"
    for ref in ("B8", "C8"):
        ws[ref].alignment = CENTER
    for r in (8, 9):
        for c in (2, 3):
            ws.cell(r, c).border = WBOX
    ws.conditional_formatting.add("C8", FormulaRule(formula=["$C$8<0"], font=Font(name=FONT, color=RED, size=15, bold=True)))

    ws["D4"] = "Winrate"
    ws["D4"].font = font(BLUE, 11, True)
    ws["D4"].alignment = Alignment(horizontal="right", vertical="center")
    ws["E4"] = "=IF($C$5+$C$6=0,0,$C$5/($C$5+$C$6))"
    ws["E4"].number_format = "0.0%"
    ws["E4"].font = font(GREEN, 11, True)
    ws["E4"].alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[10].height = 20
    ws.merge_cells("B10:I10")
    ws["B10"] = "Ky raport pasqyron rezultate të kaluara dhe nuk garanton performancë të ardhshme."
    ws["B10"].font = font(GREY, 9, italic=True, bold=True)
    ws["B10"].alignment = CENTER

    # --- tabela
    headers = ["Data", "Pair", "Setup", "R:R", "Status i pozitës", "Net", "Balance", "Shënime"]
    ws.row_dimensions[11].height = 22
    for i, h in enumerate(headers):
        c = ws.cell(11, 2 + i, h)
        c.font, c.fill, c.alignment, c.border = font(BLUE, 11, True), fill(HEAD_BG), CENTER, BOX
    for col, h in (("Z", "Brenda periudhës"), ("AA", "ln(1+Net)"), ("AB", "Kategoria"), ("AC", "cTrader ID")):
        ws[f"{col}11"] = h
        ws[f"{col}11"].font = font(GREY, 9, True)

    status_rng = f"{SETQ}!$A$6:$A$10"
    for r in range(FIRST, LAST + 1):
        ws.row_dimensions[r].height = 19
        for c in range(2, 10):
            cell = ws.cell(r, c)
            cell.border = BOX
            cell.fill = fill(INPUT_BG if c in (2, 3, 4, 5, 6, 9) else ROW_BG)
            cell.alignment = LEFT if c == 9 else CENTER
        ws.cell(r, 2).number_format = "dd mmm yy"
        ws.cell(r, 5).number_format = "0.00"
        ws.cell(r, 6).font = font(bold=True)
        ws.cell(r, 9).font = font(GREY, 9)

        ws.cell(r, 7, f'=IF($AB{r}="","",{SETQ}!$C$3*('
                      f'INDEX({SETQ}!$C$6:$C$10,MATCH($F{r},{status_rng},0))*N($E{r})'
                      f'+INDEX({SETQ}!$D$6:$D$10,MATCH($F{r},{status_rng},0))))')
        ws.cell(r, 7).number_format = "0.00%"
        ws.cell(r, 7).font = font(GREY, bold=True)
        ws.cell(r, 8, f'=IF($Z{r}=1,EXP(SUM($AA${FIRST}:$AA{r}))-1,"")')
        ws.cell(r, 8).number_format = "0.00%"
        ws.cell(r, 8).font = font(bold=True)
        # kolonat ndihmëse (të fshehura)
        ws[f"Z{r}"] = (f'=IF($AB{r}="",0,IF(ISNUMBER($B{r}),IF(AND($B{r}>=N($B$3),'
                       f'OR(NOT(ISNUMBER($D$3)),$B{r}<=$D$3)),1,0),1))')
        # rritje e përbërë si shumë logaritmesh: një rresht i fshirë nuk prish rreshtat e tjerë
        ws[f"AA{r}"] = f"=IF($Z{r}=1,LN(MAX(1E-9,1+N($G{r}))),0)"
        ws[f"AB{r}"] = f'=IF($F{r}="","",IFERROR(INDEX({SETQ}!$B$6:$B$10,MATCH($F{r},{status_rng},0)),""))'
    for col in ("Z", "AA", "AB", "AC"):
        ws.column_dimensions[col].hidden = True

    for i, (d, pair, setup, status, rr, note) in enumerate(trades):
        r = FIRST + i
        ws.cell(r, 2, date.fromisoformat(d))
        ws.cell(r, 3, pair)
        ws.cell(r, 4, setup)
        ws.cell(r, 5, rr)
        ws.cell(r, 6, status)
        if note:
            ws.cell(r, 9, note)

    # --- formatimi me kusht
    body = f"F{FIRST}:F{LAST}"
    for st, (bg, fg) in STATUS_STYLE.items():
        ws.conditional_formatting.add(body, FormulaRule(
            formula=[f'UPPER(TRIM($F{FIRST}))="{st}"'], fill=fill(bg), font=Font(name=FONT, color=fg, bold=True)))
    net = f"G{FIRST}:G{LAST}"
    ws.conditional_formatting.add(net, FormulaRule(formula=[f"AND(ISNUMBER($G{FIRST}),$G{FIRST}>0)"],
                                                   font=Font(name=FONT, color=GREEN, bold=True)))
    ws.conditional_formatting.add(net, FormulaRule(formula=[f"AND(ISNUMBER($G{FIRST}),$G{FIRST}<0)"],
                                                   font=Font(name=FONT, color=RED, bold=True)))

    # --- listat zbritëse
    dv_status = DataValidation(type="list", formula1=f"={status_rng}", allow_blank=True,
                               showErrorMessage=True, errorTitle="Status i pavlefshëm",
                               error="Zgjidh një status nga lista (ose shtoje te Cilësimet).")
    dv_setup = DataValidation(type="list", formula1=f"={SETQ}!$F$5:$F$12", allow_blank=True, showErrorMessage=False)
    dv_pair = DataValidation(type="list", formula1=f"={SETQ}!$H$5:$H$24", allow_blank=True, showErrorMessage=False)
    dv_rr = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                           showErrorMessage=True, errorTitle="R:R", error="Shkruaj një numër, p.sh. 1.4")
    dv_date = DataValidation(type="date", operator="between", formula1="36526", formula2="73051", allow_blank=True,
                             showErrorMessage=True, errorTitle="Data", error="Shkruaj një datë, p.sh. 07/01/2026")
    for dv, col in ((dv_status, "F"), (dv_setup, "D"), (dv_pair, "C"), (dv_rr, "E"), (dv_date, "B")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{LAST}")

    # --- grafikët
    chart_last = FIRST + len(trades) - 1 if demo else LAST
    ws.add_chart(winrate_chart(ws), "D5")
    mini = balance_chart(ws, chart_last, 7.6, 4.2, 1000)
    mini.y_axis.txPr = axis_text(size=700)
    ws.add_chart(mini, "F4")
    ws.add_chart(balance_chart(ws, chart_last, 18.5, 12.5, 1600), "K2")

    ws["K29"] = "Plotëso vetëm: Data, Pair, Setup, R:R, Status i pozitës dhe Shënime (qelizat me sfond më të çelët)."
    ws["K30"] = "Net, Balance, statistikat dhe grafikët llogariten vetë. Periudha ndryshohet te B3 dhe D3."
    for ref in ("K29", "K30"):
        ws[ref].font = font(GREY, 9, italic=True)
    ws.sheet_view.selection[0].activeCell = f"B{FIRST}"
    ws.sheet_view.selection[0].sqref = f"B{FIRST}"
    return ws


def build_guide(wb):
    ws = wb.create_sheet("Udhëzime")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    for r in range(1, 40):
        for c in (1, 2):
            ws.cell(r, c).fill = fill(BG)
    lines = [
        ("UDHËZIME", "title"),
        ("", None),
        ("Si ta përdorësh", "head"),
        ("1. Te fleta 'Raporti' shto një rresht për çdo trade, duke nisur nga rreshti 12.", None),
        ("2. Plotëso: Data, Pair, Setup, R:R dhe Status i pozitës (zgjidhe nga lista). Shënime është opsionale.", None),
        ("3. Net, Balance, Total Trades/WINS/LOSS/BE, Winrate, PnL dhe grafikët llogariten automatikisht – mos i shkruaj me dorë.", None),
        ("4. Periudha e raportit: ndrysho datat te B3 (nga) dhe D3 (deri). Llogariten vetëm trade-t brenda periudhës.", None),
        ("5. Fshi rreshtin shembull: zgjidh B12:I12 dhe shtyp Delete. Ka vend për 500 trade (rreshtat 12–511).", None),
        ("6. Shkruaji trade-t me radhë kronologjike – Balance llogaritet nga lart poshtë.", None),
        ("", None),
        ("Si llogaritet", "head"),
        ("• Net = Rreziku × (R:R × shumëzuesi + R fikse). Rreziku dhe shumëzuesit janë te fleta 'Cilësimet'.", None),
        ("• TP = 1 × R:R · CLOSE HALF - TP = 0.75 × R:R · CLOSE HALF - BE = 0.25 × R:R · BE = 0 · LOSS = −1R", None),
        ("• Balance rritet me interes të përbërë: (1 + Balance i mëparshëm) × (1 + Net) − 1. PnL = Balance i fundit.", None),
        ("• Winrate = WINS / (WINS + LOSS). BE nuk ndikon.", None),
        ("• Me Rrezik 23% dhe R:R 1.40 merr saktë numrat e raportit origjinal (8.05%, 24.15%, −23.00%).", None),
        ("", None),
        ("Google Sheets", "head"),
        ("Ngarkoje skedarin në Google Drive → hape me Google Sheets (ose: File → Import → Upload).", None),
        ("", None),
        ("Lidhja me cTrader", "head"),
        ("Trade-t mund të vijnë automatikisht nga cTrader: shih ctrader/UDHEZIME.md (cBot RaportiSync + Google Apps Script).", None),
        ("", None),
        ("Ngjyrat", "head"),
        ("Qelizat me sfond pak më të çelët janë për t'u plotësuar. Qelizat e errëta llogariten vetë.", None),
    ]
    for i, (text, kind) in enumerate(lines, start=2):
        c = ws.cell(i, 2, text)
        if kind == "title":
            c.font = font(size=18, bold=True)
        elif kind == "head":
            c.font = font(BLUE, 12, True)
        else:
            c.font = font(size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[i].height = 26 if kind == "title" else 18
    return ws


def main():
    demo = "--demo" in sys.argv
    if demo:
        out = sys.argv[sys.argv.index("--demo") + 1]
        trades = [(d, p, s, st, 1.4, None) for d, p, s, st in DEMO]
        risk = 0.23
    else:
        out = sys.argv[1] if len(sys.argv) > 1 else "Raporti_i_Analizave.xlsx"
        trades = [("2026-01-07", "EUR/JPY", "IMB", "CLOSE HALF - BE", 1.4, "Shembull – fshije këtë rresht")]
        risk = 0.01

    wb = Workbook()
    build_main(wb, trades, demo)
    build_settings(wb, risk)
    build_guide(wb)
    wb.save(out)
    print("U ruajt:", out)


if __name__ == "__main__":
    main()
